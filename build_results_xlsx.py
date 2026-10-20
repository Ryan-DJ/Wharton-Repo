"""Collect every backtest result into Laura_Gao_Backtest_Results.xlsx (values from the saved runs; summaries as formulas)."""
import glob, os, pickle
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from laura_strategy import Plan, default_phases, project, v3_phases

NAVY = "011F5B"; LIGHT = "EEF3FA"; YEL = "FFF2CC"
F = lambda **k: Font(name="Arial", **{"size": 10, **k})
HDR = dict(font=Font(name="Arial", size=10, bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=NAVY))
thin = Side(style="thin", color="C8CFDB")


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = HDR["font"], HDR["fill"]
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_table(ws, r0, header, rows, fmts):
    for j, h in enumerate(header, 1):
        ws.cell(row=r0, column=j, value=h)
    style_header(ws, r0, len(header))
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r0 + i, column=j, value=v)
            c.font = F(); c.border = Border(bottom=thin)
            if fmts[j - 1]:
                c.number_format = fmts[j - 1]
    return r0 + len(rows)


def widths(ws, w):
    for j, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(j)].width = x


USD, PCT, INT = '$#,##0;($#,##0);"-"', '0.0%;-0.0%;"-"', '0'


def full_plan_rows():
    """One row per (start year, version) from every saved full-plan run."""
    runs = []
    for f in sorted(glob.glob("run_*.pkl")):
        ver, yr = os.path.basename(f)[4:-4].rsplit("_", 1)
        runs.append((int(yr) + 1, {"v2": "v2 (current)", "v3": "v3 (before fixes)", "v3fix": "v3 final"}[ver], pickle.load(open(f, "rb"))["summary"], "S&P 500 point-in-time, sector names fixed"))
    have = {(y, v) for y, v, _, _ in runs}
    for f in ("us_lifecycle_1996_2002.pkl", "us_lifecycle_2003_2009.pkl"):
        for (s, mode), x in pickle.load(open(f, "rb")).items():
            y = int(s[:4]) + 1
            v = "v2 (current)" if mode == "funnel" else "Same mix, index funds"
            if (y, v) not in have:
                runs.append((y, v, x["summary"], "S&P 500 point-in-time" + (", before sector-name fix" if mode == "funnel" else "")))
    order = {"v2 (current)": 0, "v3 (before fixes)": 1, "v3 final": 2, "Same mix, index funds": 3}
    runs.sort(key=lambda r: (r[0], order[r[1]]))
    rows = []
    for y, v, sm, note in runs:
        a = sm["accounts"]; g = a.get("growth", {}); fu = a.get("funding", {}); rs = a.get("reserve", {})
        rows.append([y, f"{y}-{y + 15}", v, g.get("cagr"), g.get("max_dd"), sm["value_2031"], min(450_000, sm["value_2031"]),
                     max(0, sm["value_2031"] - 450_000), sm["funding_2033_before_gift"], sm["gift"], sm["payments_made"],
                     sm["total_shortfall"], fu.get("cagr") if sm["gift"] > 0 else None, fu.get("max_dd") if sm["gift"] > 0 else None,
                     rs.get("cagr"), rs.get("max_dd"), sm["reserve_left_2043"], sm["funding_left_2043"], sm["trading_costs"], note])
    return rows


def main(out="Laura_Gao_Backtest_Results.xlsx"):
    wb = Workbook()
    # ------------------------------------------------------------------ README
    ws = wb.active; ws.title = "README"
    lines = [
        ("Laura Gao strategy: every backtest result", True),
        ("Real-market backtests: on each date the algorithm could only buy stocks that were in the S&P 500 that day (member list since 1996).", False),
        ("A start year plays the role of Jan 2027 (year 1). Example: start 2007 = history 2007-2022 replayed as 2027-2042.", False),
        ("", False),
        ("Versions", True),
        ("v2 (current): the algorithm before the backtest revisions. Growth: 30% bonds fixed, PSR >= 0.80, drawdown >= -60%, rebalance every 6 months.", False),
        ("v3 (before fixes): first revision - crash switch (gate M1), ranked PSR, market-relative drawdown, momentum 25%, bond share chosen safest-first, T-bills allowed inside stock sleeves.", False),
        ("v3 final: v3 plus two fixes - no T-bills inside the stock sleeves (only gate M1 moves stocks to cash) and bonds added only while the target is still hit with 10 points to spare.", False),
        ("Same mix, index funds: the spreadsheet mix held in index funds (SPY / EFA / IEF), no stock picking - the yardstick.", False),
        ("", False),
        ("Sheets", True),
        ("Full plan: the whole 2027-2042 plan (contributions, 2031 split, 2033 gift, ten $50k payments) per start year and version.", False),
        ("v2 vs v3 final: side-by-side for the start years run with both (differences are formulas).", False),
        ("Phase tests: each phase on its own over its real length, every start year (v2 settings). Summary rows are formulas.", False),
        ("Spreadsheet targets: the team's bear / base / bull plan, as the code reproduces it.", False),
        ("Demo tests (biased): earlier tests on 41 stocks picked in 2026. Kept for the record; they flatter results.", False),
        ("", False),
        ("Limits", True),
        ("Yahoo has prices for 738 of 1,209 past S&P 500 members (47% of members in 1999, 97% in 2024); missing ones are mostly failed or acquired companies.", False),
        ("Measured effect: the members we have beat the real equal-weight S&P 500 fund (RSP) by about 1.5% a year (2003-2026), so stock returns here are flattered by about that much.", False),
        ("No company statements before 2021 in free data, so in these backtests gates A3-A5 (valuation, solvency, earnings quality) are skipped; all price-based gates run.", False),
        ("The foreign 20-29% holds an international index fund in backtests (no free list of past foreign index members). Bonds = IEF (Vanguard intermediate Treasury fund before 2002).", False),
        ("Trading cost 0.25% of every dollar traded. All decisions use only data available on that date.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t); c.font = F(bold=b, size=12 if i == 1 else 10)
    ws.column_dimensions["A"].width = 150
    # ------------------------------------------------------------------ Full plan
    ws = wb.create_sheet("Full plan")
    hdr = ["Start year", "History used", "Version", "Growth return / yr", "Growth worst drop", "Value at 2031 split",
           "To reserve", "To funding account", "Funding account at gift", "Facility gift", "Payments made (of 10)",
           "Payment shortfall", "Post-gift return / yr", "Post-gift worst drop", "Reserve return / yr", "Reserve worst drop",
           "Reserve left 2043", "Funding left 2043", "Trading costs", "Data"]
    rows = full_plan_rows()
    fm = [INT, None, None, PCT, PCT, USD, USD, USD, USD, USD, INT, USD, PCT, PCT, PCT, PCT, USD, USD, USD, None]
    write_table(ws, 1, hdr, rows, fm)
    for i, r in enumerate(rows, 2):
        if r[2] == "v3 final":
            for j in range(1, len(hdr) + 1):
                ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor=YEL)
    widths(ws, [9, 11, 20] + [12] * 16 + [40]); ws.freeze_panes = "D2"
    # ------------------------------------------------------------------ v2 vs v3 final (formulas)
    ws2 = wb.create_sheet("v2 vs v3 final")
    idx = {(r[0], r[2]): i + 2 for i, r in enumerate(rows)}
    years = sorted({r[0] for r in rows if (r[0], "v3 final") in idx and (r[0], "v2 (current)") in idx})
    cols = [("Value at 2031 split", "F", USD), ("Facility gift", "J", USD), ("Payments made", "K", INT),
            ("Growth worst drop", "E", PCT), ("Reserve left 2043", "Q", USD), ("Funding left 2043", "R", USD)]
    h = ["Start year"]
    for name, _, _ in cols:
        h += [f"{name}: v2", f"{name}: v3 final", f"{name}: change"]
    for j, x in enumerate(h, 1):
        ws2.cell(row=1, column=j, value=x)
    style_header(ws2, 1, len(h))
    for i, y in enumerate(years, 2):
        ws2.cell(row=i, column=1, value=y).font = F()
        a, b = idx[(y, "v2 (current)")], idx[(y, "v3 final")]
        for k, (name, col, fmt) in enumerate(cols):
            c0 = 2 + 3 * k
            for off, f_ in ((0, f"='Full plan'!{col}{a}"), (1, f"='Full plan'!{col}{b}"),
                            (2, f"={get_column_letter(c0 + 1)}{i}-{get_column_letter(c0)}{i}")):
                c = ws2.cell(row=i, column=c0 + off, value=f_); c.font = F(); c.number_format = fmt
    widths(ws2, [10] + [13] * (len(h) - 1))
    # ------------------------------------------------------------------ Phase tests (formulas for the summary)
    ws3 = wb.create_sheet("Phase tests")
    pr = pd.read_csv("us_phase_results.csv")
    order = ["growth", "pre_gift_y5", "pre_gift", "post_gift", "reserve"]
    names = {"growth": "Growth (4 yrs)", "pre_gift_y5": "Pre-gift 2031 (1 yr)", "pre_gift": "Pre-gift 2032 (1 yr)",
             "post_gift": "Post-gift (10 yrs)", "reserve": "Reserve (12 yrs)"}
    pr["o"] = pr.phase.map(order.index); pr = pr.sort_values(["o", "first_year"])
    P2 = default_phases()
    sum_hdr = ["Phase", "Tests", "Target bear", "Target base", "Target bull", "Median", "Worst", "Best", "Index mix median",
               "S&P 500 median", "Share >= bear", "Share >= base", "Worst drop", "Drop limit"]
    r0 = 1
    ws3.cell(row=r0, column=1, value="Summary (formulas over the rows below; v2 settings, real S&P 500)").font = F(bold=True, size=11)
    for j, x in enumerate(sum_hdr, 1):
        ws3.cell(row=r0 + 1, column=j, value=x)
    style_header(ws3, r0 + 1, len(sum_hdr))
    det0 = r0 + 2 + len(order) + 2
    det_hdr = ["Phase", "Start year", "Years", "Funnel return / yr", "Same mix, index funds / yr", "S&P 500 / yr",
               "Funnel worst drop", "Index mix worst drop", "Funnel worst 12 months"]
    for j, x in enumerate(det_hdr, 1):
        ws3.cell(row=det0, column=j, value=x)
    style_header(ws3, det0, len(det_hdr))
    rr = det0 + 1; span = {}
    for ph in order:
        g = pr[pr.phase == ph]; span[ph] = (rr, rr + len(g) - 1)
        for _, x in g.iterrows():
            vals = [names[ph], int(x.first_year), int(x.years), x.funnel_cagr, x.index_cagr, x.sp500_cagr, x.funnel_max_dd,
                    x.index_max_dd, x.funnel_worst_12m]
            for j, v in enumerate(vals, 1):
                c = ws3.cell(row=rr, column=j, value=v); c.font = F()
                if j >= 4: c.number_format = PCT
            rr += 1
    for i, ph in enumerate(order):
        r = r0 + 2 + i; a, b = span[ph]; p = P2[ph]
        rng = lambda col: f"{col}{a}:{col}{b}"
        vals = [names[ph], f"=COUNT({rng('D')})", p.bear, p.target, p.bull, f"=MEDIAN({rng('D')})", f"=MIN({rng('D')})",
                f"=MAX({rng('D')})", f"=MEDIAN({rng('E')})", f"=MEDIAN({rng('F')})",
                f"=COUNTIF({rng('D')},\">=\"&C{r})/B{r}", f"=COUNTIF({rng('D')},\">=\"&D{r})/B{r}", f"=MIN({rng('G')})", p.max_dd]
        for j, v in enumerate(vals, 1):
            c = ws3.cell(row=r, column=j, value=v); c.font = F(color="0000FF" if j in (3, 4, 5, 14) else "000000")
            if j >= 3: c.number_format = PCT
    widths(ws3, [22, 10, 10, 10, 12, 12, 12, 12, 14, 12, 12, 12, 12, 12])
    # ------------------------------------------------------------------ spreadsheet targets
    ws4 = wb.create_sheet("Spreadsheet targets")
    P = Plan(phases=default_phases())
    rows4 = []
    for sc in ("bear", "base", "bull"):
        x = project(P, sc)
        rows4.append([sc, x["value_2031"], x["rows"][5]["funding"], x["rows"][7]["funding_before_gift"], x["gift"], x["reserve_end"], x["funding_end"]])
    write_table(ws4, 1, ["Scenario", "Value at 2031 split", "Funding account 2031", "Funding account at gift", "Facility gift",
                         "Reserve left 2042", "Funding left 2043"], rows4, [None] + [USD] * 6)
    ws4.cell(row=6, column=1, value="Returns assumed: growth 13/15/17%, reserve 2/3/4%, pre-gift 3/5/7%, post-gift 15/18.5/21% (v3 final targets post-gift at 7/9/11%).").font = F(italic=True)
    widths(ws4, [12] + [18] * 6)
    # ------------------------------------------------------------------ demo (biased)
    ws5 = wb.create_sheet("Demo tests (biased)")
    dm = pd.read_csv("phase_results.csv")
    write_table(ws5, 1, ["Phase", "Start", "Years", "Funnel return / yr", "Index mix / yr", "Funnel worst drop", "Index worst drop"],
                [[r.phase, str(r.start), int(r.years), r.funnel_cagr, r.index_cagr, r.funnel_max_dd, r.index_max_dd] for r in dm.itertuples()],
                [None, None, INT, PCT, PCT, PCT, PCT])
    widths(ws5, [14, 12, 8, 16, 14, 16, 16])
    wb.save(out); print("saved", out, len(rows), "full-plan rows")


if __name__ == "__main__":
    main()
